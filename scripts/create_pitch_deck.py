from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


OUT = Path(__file__).resolve().parents[1] / "pitch_deck.pptx"
W, H = 13.333, 7.5
NAVY = RGBColor(8, 19, 40)
NAVY2 = RGBColor(16, 34, 64)
WHITE = RGBColor(247, 250, 252)
MUTED = RGBColor(165, 185, 209)
INK = RGBColor(19, 35, 61)
PALE = RGBColor(237, 247, 249)
LINE = RGBColor(204, 224, 231)
TEAL = RGBColor(32, 204, 176)
MINT = RGBColor(184, 255, 235)
BLUE = RGBColor(54, 130, 224)
CORAL = RGBColor(255, 132, 105)
GOLD = RGBColor(255, 202, 89)

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(W), Inches(H)
blank = prs.slide_layouts[6]


def rect(slide, x, y, w, h, fill, rounded=False, line=None):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line or fill
    shape.line.transparency = 0 if line else 100
    if rounded:
        shape.adjustments[0] = 0.12
    return shape


def add_text(slide, value, x, y, w, h, size=14, color=INK, bold=False,
             align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, margin=0.05):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(margin)
    frame.margin_top = frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = value
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return shape


def line(slide, x1, y1, x2, y2, color=TEAL, width=2):
    connector = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    connector.line.color.rgb = color
    connector.line.width = Pt(width)
    return connector


def base(dark, section, number):
    slide = prs.slides.add_slide(blank)
    rect(slide, 0, 0, W, H, NAVY if dark else WHITE)
    rect(slide, 0, 0, 0.16, H, TEAL)
    add_text(slide, section.upper(), 0.6, 0.28, 5.2, 0.24, 8, TEAL if dark else BLUE, True)
    add_text(slide, number, 12.2, 0.28, 0.5, 0.24, 9, MUTED if dark else BLUE, True, PP_ALIGN.RIGHT)
    return slide


def heading(slide, value, subtitle, dark):
    add_text(slide, value, 0.6, 0.72, 12, 0.65, 28, WHITE if dark else INK, True)
    add_text(slide, subtitle, 0.64, 1.45, 11.7, 0.4, 13, MUTED if dark else RGBColor(76, 96, 122))


def card(slide, x, y, w, h, title, body, accent, dark=True):
    rect(slide, x, y, w, h, NAVY2 if dark else PALE, True,
         RGBColor(43, 70, 104) if dark else LINE)
    rect(slide, x, y, 0.07, h, accent, True)
    add_text(slide, title, x + 0.2, y + 0.16, w - 0.35, 0.32, 14,
             WHITE if dark else INK, True)
    add_text(slide, body, x + 0.2, y + 0.58, w - 0.35, h - 0.68, 10.5,
             MUTED if dark else RGBColor(76, 96, 122))


def dot(slide, x, y, color, label=None):
    circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(0.3), Inches(0.3))
    circle.fill.solid()
    circle.fill.fore_color.rgb = color
    circle.line.fill.background()
    if label:
        add_text(slide, label, x - 0.2, y + 0.37, 0.7, 0.2, 8, MUTED, True, PP_ALIGN.CENTER)


def bullet(slide, value, x, y, w, color=WHITE, accent=TEAL, size=12):
    dot(slide, x, y + 0.09, accent)
    add_text(slide, value, x + 0.22, y, w - 0.22, 0.36, size, color)


def pill(slide, value, x, y, w, fill=TEAL, color=NAVY):
    rect(slide, x, y, w, 0.3, fill, True)
    add_text(slide, value, x, y + 0.02, w, 0.22, 8, color, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)


# 1 — opening
slide = base(True, "The new category", "01")
add_text(slide, "AGENT", 0.64, 1.1, 2.3, 0.42, 17, TEAL, True)
add_text(slide, "SUITE", 2.48, 1.1, 2.2, 0.42, 17, WHITE, True)
add_text(slide, "From company knowledge\nto autonomous, governed work", 0.62, 1.95, 7.5, 1.35, 32, WHITE, True)
add_text(slide, "AgentSuite discovers where AI can work, creates the workforce, executes the workflow, and verifies the outcome.", 0.66, 3.62, 6.9, 0.75, 16, MUTED)
pill(slide, "FROM COMPANY CONTEXT TO VERIFIED OUTCOMES", 0.66, 5.2, 3.3)
add_text(slide, "Governed digital workers for engineering and operations.", 0.66, 6.05, 6.8, 0.3, 13, MINT, True)
rect(slide, 8.45, 0.8, 3.9, 5.8, NAVY2, True, RGBColor(43, 70, 104))
labels = [("KNOWLEDGE", 9.25, 1.45, BLUE), ("WORKFORCE", 10.75, 2.55, GOLD),
          ("EXECUTION", 9.05, 4.08, CORAL), ("PROOF", 11.05, 4.92, TEAL)]
for label, x, y, color in labels:
    dot(slide, x, y, color)
    add_text(slide, label, x - 0.36, y + 0.42, 1.03, 0.2, 8, MUTED, True, PP_ALIGN.CENTER)
for x, y in [(9.4, 1.6), (10.9, 2.7), (9.2, 4.23)]:
    line(slide, x, y, 11.2, 5.07, RGBColor(56, 91, 128), 1)

# 2 — problem
slide = base(False, "The execution problem", "02")
heading(slide, "Enterprise AI has an execution problem", "More tools and more agents do not automatically create trusted work.", False)
steps = [("Too many\nsystems", BLUE), ("Too many\nworkflows", TEAL), ("Too many\npermissions", CORAL), ("No reliable\nverification", GOLD)]
for i, (label, color) in enumerate(steps):
    x = 0.72 + i * 3.02
    rect(slide, x, 2.38, 2.35, 1.18, PALE, True, LINE)
    dot(slide, x + 1.03, 2.18, color)
    add_text(slide, label, x + 0.18, 2.76, 1.99, 0.48, 15, INK, True, PP_ALIGN.CENTER)
    if i < 3:
        line(slide, x + 2.4, 2.98, x + 2.85, 2.98, RGBColor(149, 179, 192), 2)
rect(slide, 0.72, 4.4, 11.85, 1.12, NAVY, True)
add_text(slide, "AI pilots don't become trusted digital workers.", 1.0, 4.7, 11.3, 0.4, 22, WHITE, True, PP_ALIGN.CENTER)
add_text(slide, "They need context, bounded authority, durable execution, and proof of outcome.", 1.0, 5.95, 11.3, 0.3, 14, INK, True, PP_ALIGN.CENTER)

# 3 — big idea
slide = base(True, "The big idea", "03")
heading(slide, "AgentSuite builds the workforce before it executes the work", "Company knowledge becomes an operating model for governed digital workers.", True)
flow = [("Company\nKnowledge", BLUE), ("Business\nUnderstanding", TEAL), ("AI\nOrganization", GOLD),
        ("AI\nWorkforce", CORAL), ("Real Workflow\nExecution", MINT), ("Validation", TEAL)]
xs = [0.68, 2.72, 4.76, 6.8, 8.84, 10.88]
for i, ((label, color), x) in enumerate(zip(flow, xs)):
    rect(slide, x, 2.55, 1.62, 1.15, NAVY2, True, RGBColor(43, 70, 104))
    dot(slide, x + 0.66, 2.28, color)
    add_text(slide, label, x + 0.08, 2.9, 1.46, 0.52, 11.5, WHITE, True, PP_ALIGN.CENTER)
    if i < len(flow) - 1:
        line(slide, x + 1.69, 3.12, xs[i + 1] - 0.08, 3.12, TEAL, 2)
add_text(slide, "Knowledge reveals what the company does.", 0.72, 4.65, 3.4, 0.28, 14, MINT, True)
add_text(slide, "Discovery identifies repeatable work.", 4.6, 4.65, 3.4, 0.28, 14, WHITE, True)
add_text(slide, "Roles and tools become an executable workforce.", 8.48, 4.65, 4.1, 0.28, 14, WHITE, True)
rect(slide, 0.72, 5.35, 11.8, 0.8, RGBColor(22, 51, 80), True)
add_text(slide, "The Director routes intent to the right specialists — it does not perform writes.", 1.0, 5.58, 11.2, 0.28, 14, MINT, True, PP_ALIGN.CENTER)

# 4 — discovery
slide = base(False, "Grounded discovery", "04")
heading(slide, "Don’t start with agents. Start with the work.", "AI-assisted discovery turns organizational knowledge into grounded agent specifications.", False)
flow = [("Company data", BLUE), ("Capabilities", TEAL), ("Teams + roles", GOLD),
        ("Workflows", CORAL), ("Agent specs", BLUE)]
for i, ((label, color), x) in enumerate(zip(flow, [0.75, 3.25, 5.75, 8.25, 10.75])):
    rect(slide, x, 2.45, 1.7, 0.95, PALE, True, LINE)
    rect(slide, x, 2.45, 1.7, 0.08, color, True)
    add_text(slide, label, x + 0.08, 2.78, 1.54, 0.25, 12, INK, True, PP_ALIGN.CENTER)
    if i < 4:
        line(slide, x + 1.78, 2.93, x + 2.38, 2.93, RGBColor(149, 179, 192), 2)
card(slide, 0.75, 4.2, 3.72, 1.4, "Grounded", "Resources, ownership, scope and evidence come from the Knowledge Graph.", BLUE, False)
card(slide, 4.8, 4.2, 3.72, 1.4, "AI-assisted", "The system proposes meanings, capabilities and specializations; it does not invent facts.", TEAL, False)
card(slide, 8.85, 4.2, 3.72, 1.4, "Ready to activate", "AWS, Slack and Notion become visible, blocked agent stubs until their connectors are wired.", CORAL, False)

# 5 — execution examples
slide = base(True, "Real workflow execution", "05")
heading(slide, "A digital worker should complete the workflow — not just describe it", "The Director selects specialists; each specialist acts only inside its approved scope.", True)
rect(slide, 0.7, 2.22, 5.73, 3.55, NAVY2, True, RGBColor(43, 70, 104))
add_text(slide, "MEETING REQUEST", 1.0, 2.52, 2.0, 0.22, 9, TEAL, True)
meeting = [("Customer asks to book a meeting", BLUE), ("Director routes to Gmail + Calendar", GOLD),
           ("Calendar creates the event", TEAL), ("Gmail replies with the meeting link", MINT)]
for i, (label, color) in enumerate(meeting):
    y = 3.0 + i * 0.6
    dot(slide, 1.05, y + 0.04, color)
    add_text(slide, label, 1.48, y, 4.35, 0.3, 12, WHITE, i == 1)
    if i < 3:
        line(slide, 1.2, y + 0.34, 1.2, y + 0.58, RGBColor(56, 91, 128), 1)
rect(slide, 6.85, 2.22, 5.73, 3.55, NAVY2, True, RGBColor(43, 70, 104))
add_text(slide, "CUSTOMER COMPLAINT", 7.15, 2.52, 2.35, 0.22, 9, CORAL, True)
complaint = [("Customer reports a feature issue", BLUE), ("Gmail agent extracts the context", GOLD),
             ("Director routes to GitHub specialist", TEAL), ("Issue created; customer gets a reply", MINT)]
for i, (label, color) in enumerate(complaint):
    y = 3.0 + i * 0.6
    dot(slide, 7.2, y + 0.04, color)
    add_text(slide, label, 7.63, y, 4.35, 0.3, 12, WHITE, i == 2)
    if i < 3:
        line(slide, 7.35, y + 0.34, 7.35, y + 0.58, RGBColor(56, 91, 128), 1)
add_text(slide, "From company context to verified outcomes.", 0.72, 6.28, 11.85, 0.3, 16, MINT, True, PP_ALIGN.CENTER)

# 6 — trust
slide = base(False, "Trust architecture", "06")
heading(slide, "AI decides. Policy constrains. Evidence proves.", "Three distinct responsibilities keep autonomy useful and safe.", False)
card(slide, 0.72, 2.35, 3.72, 2.32, "AI / Director", "Understands intent, identifies capability, and selects a small set of specialist agents. No write tools.", BLUE, False)
card(slide, 4.8, 2.35, 3.72, 2.32, "Guardrails / Policy", "Determines what the agent is allowed to do: scope, tool access, allow, deny, or human approval.", CORAL, False)
card(slide, 8.88, 2.35, 3.72, 2.32, "Validation", "Re-queries the live system and checks whether the intended external outcome actually happened.", TEAL, False)
add_text(slide, "Can the agent do this?", 1.05, 5.35, 3.1, 0.28, 14, BLUE, True, PP_ALIGN.CENTER)
add_text(slide, "Did it actually happen correctly?", 8.98, 5.35, 3.5, 0.28, 14, TEAL, True, PP_ALIGN.CENTER)
line(slide, 4.15, 5.48, 8.7, 5.48, RGBColor(149, 179, 192), 2)
add_text(slide, "Permission is not proof.", 5.05, 5.88, 3.3, 0.3, 16, INK, True, PP_ALIGN.CENTER)

# 7 — validation
slide = base(True, "Evidence-based outcomes", "07")
heading(slide, "Successful execution is not the same as successful outcome", "Agent claims are not enough; validation is based on observable evidence.", True)
steps = [("EXECUTE", BLUE), ("TOOL EVENTS", TEAL), ("GUARDRAIL\nDECISION", GOLD),
         ("OBSERVED\nOUTCOME", CORAL), ("INDEPENDENT\nVALIDATION", MINT)]
for i, ((label, color), x) in enumerate(zip(steps, [0.72, 3.1, 5.48, 7.86, 10.24])):
    rect(slide, x, 2.5, 1.82, 0.92, NAVY2, True, RGBColor(43, 70, 104))
    rect(slide, x, 2.5, 1.82, 0.07, color, True)
    add_text(slide, label, x + 0.1, 2.82, 1.62, 0.35, 10.5, WHITE, True, PP_ALIGN.CENTER)
    if i < 4:
        line(slide, x + 1.9, 2.96, x + 2.27, 2.96, TEAL, 2)
rect(slide, 0.72, 4.45, 5.45, 1.05, RGBColor(24, 53, 76), True)
add_text(slide, "Guardrail: ALLOW send_email", 1.0, 4.73, 4.9, 0.25, 14, MINT, True)
rect(slide, 6.52, 4.45, 6.05, 1.05, RGBColor(68, 39, 43), True)
add_text(slide, "Validation: FAIL — external outcome not confirmed", 6.8, 4.73, 5.5, 0.25, 13, CORAL, True)
add_text(slide, "Evidence outcomes: PASS  •  FAIL  •  NO_EVIDENCE", 0.72, 6.25, 11.85, 0.3, 15, WHITE, True, PP_ALIGN.CENTER)

# 8 — controlled learning
slide = base(False, "Controlled learning loop", "08")
heading(slide, "Every validated execution creates feedback for improvement", "Future capability: improve behavior and validation through evaluation — never through self-modifying permissions.", False)
loop = [("EXECUTE", BLUE), ("EVIDENCE", TEAL), ("VALIDATE", GOLD), ("HUMAN\nREVIEW", CORAL),
        ("EVALUATION\nDATA", BLUE), ("PROPOSE\nIMPROVEMENT", TEAL), ("TEST /\nAPPROVE", INK)]
for i, ((label, color), x) in enumerate(zip(loop, [0.62, 2.42, 4.22, 6.02, 7.82, 9.62, 11.42])):
    rect(slide, x, 2.52, 1.36, 0.95, PALE, True, LINE)
    rect(slide, x, 2.52, 1.36, 0.07, color, True)
    add_text(slide, label, x + 0.05, 2.83, 1.26, 0.34, 9.5, INK, True, PP_ALIGN.CENTER)
    if i < 6:
        line(slide, x + 1.43, 3.0, x + 1.73, 3.0, RGBColor(149, 179, 192), 2)
line(slide, 12.1, 3.62, 12.1, 4.35, TEAL, 2)
line(slide, 12.1, 4.35, 1.3, 4.35, TEAL, 2)
line(slide, 1.3, 4.35, 1.3, 3.62, TEAL, 2)
add_text(slide, "Failed executions become regression cases. Human reviews become labeled feedback. Improvements are promoted only after testing.", 1.0, 4.8, 11.2, 0.55, 15, INK, True, PP_ALIGN.CENTER)
rect(slide, 2.65, 5.8, 8.0, 0.62, NAVY, True)
add_text(slide, "Policies and permissions remain operator-controlled.", 2.9, 5.98, 7.5, 0.23, 13, MINT, True, PP_ALIGN.CENTER)

# 9 — connected capabilities
slide = base(False, "Category definition", "09")
heading(slide, "What AgentSuite connects", "The differentiation is the loop — not a single model, tool, or connector.", False)
headers = [("Traditional approach", 0.8, PALE, INK), ("AgentSuite", 6.95, NAVY, WHITE)]
for label, x, fill, color in headers:
    rect(slide, x, 2.16, 5.58, 0.56, fill, True)
    add_text(slide, label, x, 2.33, 5.58, 0.22, 14, color, True, PP_ALIGN.CENTER)
rows = [("Company context", "Knowledge-grounded"), ("Manual agent creation", "Grounded work discovery"),
        ("Individual agents", "AI workforce"), ("Tool calls", "Governed execution"),
        ("Claimed success", "Verified outcome"), ("Static system", "Controlled improvement")]
for i, (left, right) in enumerate(rows):
    y = 2.9 + i * 0.53
    rect(slide, 0.8, y, 5.58, 0.42, WHITE, True, LINE)
    rect(slide, 6.95, y, 5.58, 0.42, NAVY2, True, RGBColor(43, 70, 104))
    add_text(slide, left, 0.98, y + 0.11, 5.2, 0.18, 10.5, RGBColor(76, 96, 122), align=PP_ALIGN.CENTER)
    add_text(slide, right, 7.13, y + 0.11, 5.2, 0.18, 10.5, WHITE, True, align=PP_ALIGN.CENTER)
add_text(slide, "From discovery to execution to proof — one connected operating layer.", 0.8, 6.35, 11.7, 0.3, 15, INK, True, PP_ALIGN.CENTER)

# 10 — close / pilot
slide = base(True, "The path forward", "10")
heading(slide, "Start with one department. Build the workforce around real work.", "Move from AI experiments to trusted digital workers.", True)
steps = [("1", "Department", BLUE), ("2", "1–3 workflows", TEAL), ("3", "Governed agents", GOLD),
         ("4", "Measured outcomes", CORAL), ("5", "Validated feedback", MINT)]
for i, (number, label, color) in enumerate(steps):
    x = 0.78 + i * 2.45
    dot(slide, x + 0.68, 2.42, color)
    add_text(slide, number, x + 0.68, 2.49, 0.3, 0.16, 9, NAVY, True, PP_ALIGN.CENTER)
    add_text(slide, label, x, 3.0, 1.66, 0.26, 12, WHITE, True, PP_ALIGN.CENTER)
    if i < 4:
        line(slide, x + 1.82, 2.57, x + 2.28, 2.57, TEAL, 2)
rect(slide, 1.0, 4.4, 11.3, 1.25, TEAL, True)
add_text(slide, "Connect the work. Govern the action. Verify the outcome.", 1.3, 4.74, 10.7, 0.35, 22, NAVY, True, PP_ALIGN.CENTER)
add_text(slide, "AgentSuite: from company context to verified outcomes.", 0.72, 6.25, 11.85, 0.3, 16, MINT, True, PP_ALIGN.CENTER)

prs.core_properties.title = "AgentSuite: From Company Knowledge to Governed Work"
prs.core_properties.subject = "Enterprise pitch deck"
prs.core_properties.author = "AgentSuite"
prs.core_properties.keywords = "knowledge graph, digital workers, governance, validation, Temporal"
prs.save(OUT)
print(f"Created {OUT}")
